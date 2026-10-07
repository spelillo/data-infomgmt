"""Builds test/fixtures/excel-structures.xlsx with openpyxl.

openpyxl writes the same OOXML parts Excel does for tables, data validation,
conditional formatting and charts, so this file stands in for a student's
upload made in real Excel. Re-run with:  python3 test/fixtures/make_fixtures.py
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule, Rule
from openpyxl.styles import PatternFill, Font
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

wb = Workbook()
ws = wb.active
ws.title = "Sales"
rows = [["Region", "Product", "Units", "Revenue"],
        ["East", "Pens", 10, 25.5], ["West", "Paper", 4, 12.0],
        ["East", "Paper", 7, 21.0], ["North", "Pens", 12, 30.6],
        ["West", "Pens", 3, 7.65]]
for r in rows:
    ws.append(r)
ws["F1"] = "Total"
ws["F2"] = "=SUM(D2:D6)"
table = Table(displayName="SalesTbl", ref="A1:D6")
table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium7", showRowStripes=True)
ws.add_table(table)

dv = DataValidation(type="list", formula1='"East,West,North,South"', allow_blank=True)
dv.add("A2:A50")
ws.add_data_validation(dv)
dv2 = DataValidation(type="whole", operator="between", formula1="0", formula2="100")
dv2.add("C2:C50")
ws.add_data_validation(dv2)

lists = wb.create_sheet("Lists")
for i, v in enumerate(["Pens", "Paper", "Ink"], start=1):
    lists.cell(row=i, column=1, value=v)
dv3 = DataValidation(type="list", formula1="Lists!$A$1:$A$3")
dv3.add("B2:B50")
ws.add_data_validation(dv3)

green = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
ws.conditional_formatting.add("D2:D6", CellIsRule(operator="greaterThan", formula=["20"], fill=green, font=Font(color="006100")))
red = DifferentialStyle(fill=PatternFill(bgColor="FFC7CE"), font=Font(color="9C0006"))
ws.conditional_formatting.add("B2:B6", Rule(type="duplicateValues", dxf=red))
ws.conditional_formatting.add("A2:A6", CellIsRule(operator="equal", formula=['"East"'], fill=PatternFill(bgColor="FFEB9C", fill_type="solid")))

summary = wb.create_sheet("Summary")
for r in [["Region", "Revenue"], ["East", 46.5], ["West", 19.65], ["North", 30.6]]:
    summary.append(r)
bar = BarChart()
bar.type = "col"
bar.title = "East leads revenue"
bar.y_axis.scaling.min = 0
bar.add_data(Reference(summary, min_col=2, min_row=1, max_row=4), titles_from_data=True)
bar.set_categories(Reference(summary, min_col=1, min_row=2, max_row=4))
summary.add_chart(bar, "D2")
line = LineChart()
line.add_data(Reference(ws, min_col=4, min_row=1, max_row=6), titles_from_data=True)
line.set_categories(Reference(ws, min_col=1, min_row=2, max_row=6))
summary.add_chart(line, "D20")

out = Path(__file__).with_name("excel-structures.xlsx")
wb.save(out)
print("wrote", out)
